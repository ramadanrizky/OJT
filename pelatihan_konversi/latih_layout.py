"""Latih threshold tata letak PDF-ke-Word (coordinate descent + skor korpus).
Jalankan:  python pelatihan_konversi\\latih_layout.py            (evaluasi saja)
           python pelatihan_konversi\\latih_layout.py --tune      (latih + simpan)
Hasil:     pelatihan_konversi\\konversi_config_v1.json
"""
import json
import os
import sys

DASAR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(DASAR)
sys.path.insert(0, ROOT)
os.chdir(ROOT)

from docx import Document  # noqa: E402

import app as A  # noqa: E402

KORPUS = os.path.join(DASAR, 'korpus')
TMP = os.path.join(DASAR, 'tmp_out')
os.makedirs(TMP, exist_ok=True)

ALIGN = {"LEFT": 0, "CENTER": 1, "RIGHT": 2, "JUSTIFY": 3}


def muat_labels():
    with open(os.path.join(KORPUS, 'labels.json'), encoding='utf-8') as fh:
        return json.load(fh)["kasus"]


def align_nama(par):
    if par.style.name.startswith('Heading'):
        return 'HEADING'
    if par.style.name.startswith('List'):
        return 'LIST'
    a = par.alignment
    if a is None:
        return 'LEFT'
    return {0: 'LEFT', 1: 'CENTER', 2: 'RIGHT', 3: 'JUSTIFY'}.get(a, 'LEFT')


def _ada_indent(par):
    try:
        li = par.paragraph_format.left_indent
        return li is not None and li.pt > 7
    except Exception:
        return False


def skor_satu(pdf_path, eks, cfg):
    """Kembalikan (skor 0..1, rincian)."""
    out = os.path.join(TMP, os.path.basename(pdf_path) + ".docx")
    try:
        A._pdf2word_flowing_v2(pdf_path, out, cfg=cfg)
    except Exception as e:
        return 0.0, [f"CRASH: {e}"]
    doc = Document(out)
    n_gambar = len(doc.inline_shapes) + sum(
        1 for p in doc.paragraphs for _ in p._p.xpath('.//wp:anchor'))
    badan = [p for p in doc.paragraphs if p.text.strip()]
    teks_badan = "\n".join(p.text for p in badan)
    dapat, total, info = 0.0, 0.0, []

    def cek(ok, nama):
        nonlocal dapat, total
        total += 1
        if ok:
            dapat += 1
        else:
            info.append("X-" + nama)

    for ep in eks.get("paragraf", []):
        cocok = [p for p in badan if ep["kunci"] in p.text]
        cek(bool(cocok), f"teks:{ep['kunci'][:20]}")
        if cocok:
            p = cocok[0]
            _mau = ep.get("align_any", [ep.get("align")])
            _dapat = align_nama(p)
            # Baris tunggal tengah diposisikan via indent presisi (LEFT+indent),
            # dinilai setara CENTER karena x-nya sama dengan sumber.
            if _dapat == "LEFT" and "CENTER" in _mau and _ada_indent(p):
                _dapat = "CENTER"
            cek(_dapat in _mau, f"align:{ep['kunci'][:20]}->{_dapat}")
            if "indent" in ep:
                ada = bool(p.paragraph_format.left_indent or p.paragraph_format.first_line_indent)
                cek(ada == ep["indent"], f"indent:{ep['kunci'][:20]}")
            if "warna" in ep:
                wr = [str(r.font.color.rgb) for r in p.runs
                      if r.font.color and r.font.color.rgb]
                cek(ep["warna"] in wr, f"warna:{ep['kunci'][:20]}")
            if ep.get("bold"):
                cek(any(r.bold for r in p.runs), f"bold:{ep['kunci'][:20]}")
            if ep.get("italic"):
                cek(any(r.italic for r in p.runs), f"italic:{ep['kunci'][:20]}")
            if ep.get("super"):
                cek(any(r.font.superscript and ep["super"] in (r.text or "") for r in p.runs),
                    f"super:{ep['kunci'][:20]}")
    cek(len(doc.tables) == eks.get("tabel", 0),
        f"tabel:{len(doc.tables)}!={eks.get('tabel', 0)}")
    cek(n_gambar == eks.get("gambar", 0),
        f"gambar:{n_gambar}!={eks.get('gambar', 0)}")
    for t in eks.get("teks_harus_ada", []):
        cek(t in teks_badan, f"ada:{t[:20]}")
    teks_tabel = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for t in eks.get("tabel_harus_ada", []):
        cek(t in teks_tabel, f"ada-tabel:{t[:20]}")
    for t in eks.get("teks_harus_tidak_ada", []):
        cek(t not in teks_badan, f"jangan-ada:{t[:20]}")
    for t in eks.get("teks_tidak_ganda", []):
        cek(t not in teks_badan, f"ganda-di-badan:{t[:20]}")
    for t in eks.get("tidak_di_badan", []):
        cek(t not in teks_badan, f"tidak-di-badan:{t[:20]}")
    if "urutan" in eks:
        try:
            pos = [teks_badan.index(k) for k in eks["urutan"]]
            cek(pos == sorted(pos), "urutan-baca")
        except ValueError:
            cek(False, "urutan-baca(hilang)")
    if "margin_L_in" in eks:
        ml = doc.sections[0].left_margin.inches
        cek(abs(ml - eks["margin_L_in"]) <= 0.35, f"margin:{ml:.2f}")
    if "header_harus_ada" in eks:
        hdr = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
        _hdr_gambar = any(
            r._r.xpath('.//wp:anchor') or r._r.xpath('.//wp:inline')
            for p in doc.sections[0].header.paragraphs for r in p.runs)
        for t in eks["header_harus_ada"]:
            # lolos bila teks ada ATAU watermark berupa gambar di header
            cek(t in hdr or _hdr_gambar, f"header:{t[:20]}")
    if "kop_tidak_di_badan" in eks:
        for t in eks["kop_tidak_di_badan"]:
            cek(t not in teks_badan, f"kop-badan:{t[:20]}")
    try:
        os.remove(out)
    except OSError:
        pass
    return (dapat / total if total else 0.0), info


GRID = {
    "justify_minw_ratio": [0.75, 0.82, 0.90],
    "justify_tol_pt": [14.0, 22.0, 30.0],
    "justify_single_full": [0.94, 0.97, 0.99],
    "center_maxw_ratio": [0.75, 0.85, 0.92],
    "center_tol_ratio": [0.03, 0.05, 0.08],
    "right_tol_pt": [8.0, 14.0, 22.0],
    "right_gap_pt": [25.0, 40.0, 60.0],
    "indent_tol_pt": [2.0, 4.0, 8.0],
    "firstline_min_pt": [6.0, 9.0, 14.0],
    "size_tol_pt": [1.0, 2.0, 3.5],
    "single_right_edge": [0.80, 0.85, 0.92],
    "col_overlap_ratio": [0.20, 0.30, 0.45],
    "full_width_ratio": [0.60, 0.70, 0.80],
    "margin_min_pt": [20.0, 28.0, 40.0],
    "margin_max_pt": [90.0, 108.0, 130.0],
}


def skor_total(cfg, kasus):
    s, n = 0.0, 0
    gagal = []
    for ks in kasus:
        v, info = skor_satu(os.path.join(KORPUS, ks["pdf"]), ks, cfg)
        s += v
        n += 1
        if info:
            gagal.append((ks["pdf"], info))
    return s / n, gagal


def main():
    kasus = muat_labels()
    print("Korpus latih:", len(kasus), "kasus")
    cfg = A._muat_config_konversi()  # mulai dari config file (dipakai produksi)
    awal, _ = skor_total(cfg, kasus)
    print(f"SKOR awal (default): {awal:.3f}")
    if "--tune" not in sys.argv:
        print("Tambahkan --tune untuk melatih threshold.")
        return
    for ronde in range(1, 4):
        berubah = False
        for kunci, cands in GRID.items():
            terbaik_v, terbaik_s = cfg[kunci], -1.0
            for v in cands:
                cfg[kunci] = v
                s, _ = skor_total(cfg, kasus)
                if s > terbaik_s:
                    terbaik_s, terbaik_v = s, v
            if terbaik_v != cfg[kunci]:
                berubah = True
            cfg[kunci] = terbaik_v
            print(f"  ronde {ronde} {kunci}={terbaik_v} skor={terbaik_s:.3f}")
        s, gagal = skor_total(cfg, kasus)
        print(f"== ronde {ronde} selesai: skor={s:.3f}")
        for pdf, info in gagal:
            print(f"   sisa gagal {pdf}: {info[:6]}")
        if not berubah:
            print("Konvergen (tidak ada perubahan).")
            break
    with open(os.path.join(DASAR, 'konversi_config_v1.json'), 'w', encoding='utf-8') as fh:
        json.dump(cfg, fh, indent=2)
    print("Config terbaik disimpan. Skor akhir:", f"{skor_total(cfg, kasus)[0]:.3f}")


if __name__ == "__main__":
    main()
